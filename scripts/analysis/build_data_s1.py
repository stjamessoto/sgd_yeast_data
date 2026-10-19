"""
build_data_s1.py — Data S1 for the Science submission: one Excel workbook with
the numbers behind Figs. 2-4 plus the per-family π estimates.

Each figure tab is computed by executing the *computation* half of the script
that draws that figure (everything above its plotting section), so the
workbook cannot drift from the published PNGs:

  Fig. 2  scripts/analysis/_generate_doc.py
  Fig. 3  assets/directional_motifs/tf_regulates_tf_cascade_topology_groups.py
  Fig. 4  assets/directional_motifs/regulatory_link_parsimony_tree.py

No figure files are written.

Run: python scripts/analysis/build_data_s1.py
Out: supplementary/data_S1.xlsx
"""

import io
import math
import sys
import contextlib
from itertools import combinations as icombs
from pathlib import Path

REPO_ROOT = Path(__file__).parent.parent.parent
sys.path.insert(0, str(REPO_ROOT))

import numpy as np
import pandas as pd
from scipy.stats import norm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from model.scruse_math import (
    expected_partial, variance_binary, second_moment_binary, estimate_pi_hat, allocate_pi,
)
from model.data_loader import load_transcription_factors, load_inheritance_vectors
from model.consensus_loader import get_pi_consensus_factor

OUT_PATH = REPO_ROOT / "supplementary" / "data_S1.xlsx"
FDR_ALPHA = 0.05


def run_compute_half(rel_path, stop_marker):
    """Exec a figure script up to (not including) its plotting section."""
    path = REPO_ROOT / rel_path
    src = path.read_text(encoding="utf-8")
    cut = src.index(stop_marker)
    g = {"__file__": str(path), "__name__": "data_s1_" + path.stem}
    sys.path.insert(0, str(path.parent))
    with contextlib.redirect_stdout(io.StringIO()):
        exec(compile(src[:cut], str(path), "exec"), g)
    return g


# ═══════════════════════════════════════════════════════════════════════════
# Fig. 2 — motif significance (70 combinations)
# ═══════════════════════════════════════════════════════════════════════════
print("Fig. 2: running _generate_doc.py computation...")
g2 = run_compute_half("scripts/analysis/_generate_doc.py", "# ── figure")
fam_list, M, N = g2["fam_list"], g2["M"], g2["N"]
fam_pi = g2["_fam_pi"]                     # Method 1, normalized (Eq. 3)
fam_by_label = {"F%d" % f["idx"]: f for f in fam_list}

fig2_rows = []
for k, res in ((3, g2["res3"]), (4, g2["res4"])):
    for r in res:
        fams = r["label"].split("+")
        pi_vec = [fam_pi[fam_by_label[x]["idx"]] for x in fams]
        pi_hat = sum(pi_vec)
        e = expected_partial(pi_hat, M, N)
        var = variance_binary(pi_vec, M, N)
        z = (r["obs"] - e) / math.sqrt(var)
        assert abs(z - r["z"]) < 1e-3 * max(1.0, abs(r["z"])), (r["label"], z, r["z"])
        fig2_rows.append({
            "label": r["label"], "fams": fams, "k": k, "obs": r["obs"],
            "pi_hat": pi_hat, "e": e, "var": var, "z": z,
            "bh_sig_pipeline": r["bh_sig"],
        })

# Unrounded two-sided p-values and BH-adjusted q-values. (The figure pipeline
# rounds p to 6 dp before BH; significance calls are identical either way —
# checked below.)
zs = np.array([r["z"] for r in fig2_rows])
log10_p = (np.log(2.0) + norm.logsf(np.abs(zs))) / np.log(10.0)
p = 2.0 * norm.sf(np.abs(zs))
n_tests = len(p)
order = np.argsort(p)
q_sorted = np.minimum.accumulate((p[order] * n_tests / np.arange(1, n_tests + 1))[::-1])[::-1]
q = np.empty(n_tests)
q[order] = np.minimum(q_sorted, 1.0)
for r, pi_, qi, lp in zip(fig2_rows, p, q, log10_p):
    r["p"], r["q"], r["log10_p"] = float(pi_), float(qi), float(lp)
    assert (qi < FDR_ALPHA) == r["bh_sig_pipeline"], r["label"]
assert n_tests == 70

# ═══════════════════════════════════════════════════════════════════════════
# Fig. 3 — five-family TF-regulates-TF cascades (21 combinations)
# ═══════════════════════════════════════════════════════════════════════════
print("Fig. 3: running tf_regulates_tf_cascade_topology_groups.py computation...")
g3 = run_compute_half(
    "assets/directional_motifs/tf_regulates_tf_cascade_topology_groups.py",
    "# ── grid figure",
)
fig3_combos = sorted(g3["combos"], key=lambda c: c["num"])
assert len(fig3_combos) == 21

# ═══════════════════════════════════════════════════════════════════════════
# Fig. 4 — regulatory-link parsimony tree
# ═══════════════════════════════════════════════════════════════════════════
print("Fig. 4: running regulatory_link_parsimony_tree.py computation...")
g4 = run_compute_half(
    "assets/directional_motifs/regulatory_link_parsimony_tree.py",
    "# ── 6. simple phylogram layout",
)
Gr, root_node = g4["Gr"], g4["root_node"]
TAXA = g4["TAXA"]


def to_newick(G, node, parent=None):
    kids = [nb for nb in G.neighbors(node) if nb != parent]
    name = "F%d" % node if node in TAXA else ""
    if not kids:
        return name
    inner = ",".join(
        "%s:%g" % (to_newick(G, ch, node), G[node][ch]["weight"]) for ch in kids
    )
    return "(%s)%s" % (inner, name)


newick = to_newick(Gr, root_node) + ";"
fig4_score = float(g4["best_score"])
n_tied = len(g4["best_trees"])
raw, all_targets = g4["raw"], g4["all_targets"]

_pi4_names = pd.read_csv(
    REPO_ROOT / "y1000plus_data" / "processed" / "pi4_snp_binding_sites.csv",
    usecols=["target_gene_id", "target_gene_name"],
).drop_duplicates("target_gene_id").set_index("target_gene_id")["target_gene_name"]

# ═══════════════════════════════════════════════════════════════════════════
# π estimates — seven estimators, per family
# Per-TF estimates are aggregated to a family as the mean over member TFs
# that have a value (the same two-level mean as Method 1's Eq. 1).
# ═══════════════════════════════════════════════════════════════════════════
print("Pi estimates: aggregating per-TF estimators to families...")
tfs = load_transcription_factors()
tf_prior = dict(zip(tfs["gene_name"].str.upper(), tfs["pi_prior"]))
ivec = load_inheritance_vectors()
m3_const = float(ivec["pi_snp"].mean())

pi3_df = pd.read_csv(REPO_ROOT / "y1000plus_data" / "processed" / "pi3_tfbs_conservation.csv")
pi3_by_tf = pi3_df.dropna(subset=["pi3_estimate"]).groupby(
    pi3_df["tf_name"].str.upper())["pi3_estimate"].mean()
pi2_df = pd.read_csv(REPO_ROOT / "y1000plus_data" / "processed" / "pi2_sequence_homology.csv")
pi2_long = pd.concat([
    pi2_df[["gene1_name", "pi2_estimate"]].rename(columns={"gene1_name": "tf"}),
    pi2_df[["gene2_name", "pi2_estimate"]].rename(columns={"gene2_name": "tf"}),
]).dropna()
pi2_by_tf = pi2_long.groupby(pi2_long["tf"].str.upper())["pi2_estimate"].mean()
pi4_df = pd.read_csv(REPO_ROOT / "y1000plus_data" / "processed" / "pi4_snp_binding_sites.csv")
pi4_by_tf = pi4_df.dropna(subset=["pi4_estimate"]).groupby(
    pi4_df["tf_name"].str.upper())["pi4_estimate"].mean()


def fam_mean(members, lookup):
    vals = [float(lookup[t]) for t in members if t in lookup]
    return (float(np.mean(vals)) if vals else None), len(vals)


pi_rows = []
for f in fam_list:
    mem = f["member_tfs"]
    m1_raw, n1 = fam_mean(mem, tf_prior)
    m4_lookup = {t: min(1.0, tf_prior[t] * get_pi_consensus_factor(t)) for t in mem if t in tf_prior}
    m4, n4 = fam_mean(mem, m4_lookup)
    m5, n5 = fam_mean(mem, pi3_by_tf)
    m6, n6 = fam_mean(mem, pi2_by_tf)
    m7, n7 = fam_mean(mem, pi4_by_tf)
    pi_rows.append({
        "fam": "F%d" % f["idx"], "go": f["go_id"], "size": f["size"], "ev": f["ev"],
        "m1_norm": fam_pi[f["idx"]], "m1": m1_raw, "m3": m3_const,
        "m4": m4, "m5": m5, "m6": m6, "m7": m7,
        "n": (n1, n4, n5, n6, n7),
    })

# Method 2 identifies only π̂ = Σπᵢ per motif; report it per combination.
m2_rows = []
for r in fig2_rows:
    ph = estimate_pi_hat(float(r["obs"]), M, N)
    w = [fam_by_label[x]["ev"] for x in r["fams"]]
    alloc = allocate_pi(ph, w) if ph is not None else [None] * len(w)
    m2_rows.append((r["label"], r["obs"], ph, dict(zip(r["fams"], alloc))))

# ═══════════════════════════════════════════════════════════════════════════
# Workbook
# ═══════════════════════════════════════════════════════════════════════════
FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=12, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
F_LABEL = Font(name=FONT, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="2C3E50")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def title_block(ws, title, notes):
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    row = 2
    for note in notes:
        ws.cell(row=row, column=1, value=note).font = F_NOTE
        row += 1
    return row + 1


def header(ws, row, names, widths=None, freeze=True):
    for j, name in enumerate(names, start=1):
        c = ws.cell(row=row, column=j, value=name)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, row, col, value, fmt=None, font=F_BODY):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    return c


# ── README ────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "README"
r0 = title_block(ws, "Data S1. Numbers underlying Figs. 2-4 and the π estimates", [])
readme = [
    ("Tab", "Contents"),
    ("Fig 2", "All 70 three- and four-family subnetwork motif tests (35 with k = 3, 35 with k = 4): "
              "families, motif size, observed count (Cartesian product of family sizes), expected count E[|M(n)|] (Theorem 4), "
              "variance bound (Binary Inheritance, Corollary 16), Z-score, Benjamini-Hochberg "
              "FDR-adjusted p-value, significance at FDR < 0.05."),
    ("Fig 3", "All 21 five-family TF-regulates-TF cascades: summary (edge count, mutual-pair "
              "count, isolated nodes, topology group) and full directed edge list."),
    ("Fig 4", "Maximum-parsimony tree: score, Newick string with branch lengths, and the binary "
              "regulatory-link character matrix (one row per π4-covered target gene)."),
    ("Pi estimates", "Per-family π from each of the seven estimators (M1-M7), plus Method 2's "
                     "per-motif π̂."),
    ("", ""),
    ("Model parameters", "m = %d GO-Process TF families, n = %d TF-family memberships, "
                         "d = n − m = %d duplication events." % (M, N, N - M)),
    ("Families", "; ".join("F%d = %s (%s, %d TFs)" % (f["idx"], f["go_id"], f["name"], f["size"])
                           for f in fam_list)),
    ("Data sources", "SGD (sgd_yeast_data/), JASPAR 2024 / Y1000+ π4 binding sites "
                     "(y1000plus_data/processed/pi4_snp_binding_sites.csv), π3 and π2 tables in "
                     "the same folder."),
    ("Generated by", "scripts/analysis/build_data_s1.py (Zenodo archive, ref. 21). Each figure tab "
                     "reruns the computation of the script that draws that figure."),
]
for i, (a, b) in enumerate(readme):
    ca = put(ws, r0 + i, 1, a, font=F_LABEL if i else F_HEAD)
    cb = put(ws, r0 + i, 2, b, font=F_BODY if i else F_HEAD)
    cb.alignment = WRAP
    if i == 0:
        ca.fill = cb.fill = FILL_HEAD
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 110

# ── Fig 2 ─────────────────────────────────────────────────────────────────
ws = wb.create_sheet("Fig 2")
r0 = title_block(ws, "Fig. 2. Significance of all 70 three- and four-family subnetwork motifs", [
    "Null model: Partial Duplication (12), m = %d, n = %d. π per family from Method 1 "
    "(evidence score normalized by the sum over all 7 families; see Pi estimates tab)." % (M, N),
    "Observed count: Cartesian product of family sizes c₁ × c₂ × … × cₖ, the number of realizable "
    "one-gene-per-family motif instances (empirical analogue of E[|M(n)|], Theorem 4).",
    "Z = (observed − expected) / √(variance bound). Two-sided p from the normal approximation; "
    "q = Benjamini-Hochberg adjusted p across all 70 tests. p and q below ~1e-308 underflow to 0; "
    "log10(p) is given exactly.",
])
cols = ["Combination", "Families", "Motif size (k)", "Family sizes", "Observed count",
        "π̂ = Σπᵢ", "Expected count E[|M(n)|]", "Variance bound (Cor. 16)", "Z-score",
        "log10(p), two-sided", "p, two-sided", "FDR-adjusted p (BH q)", "Significant (q < 0.05)"]
header(ws, r0, cols, [16, 22, 9, 16, 14, 10, 14, 16, 13, 13, 12, 14, 13])
for i, r in enumerate(fig2_rows):
    row = r0 + 1 + i
    put(ws, row, 1, r["label"])
    put(ws, row, 2, ", ".join(r["fams"]))
    put(ws, row, 3, r["k"])
    put(ws, row, 4, " × ".join(str(fam_by_label[x]["size"]) for x in r["fams"]))
    put(ws, row, 5, r["obs"], "#,##0")
    put(ws, row, 6, r["pi_hat"], "0.0000")
    put(ws, row, 7, r["e"], "0.0000")
    put(ws, row, 8, r["var"], "0.0000")
    put(ws, row, 9, "=(E{0}-G{0})/SQRT(H{0})".format(row), "#,##0.000")
    put(ws, row, 10, r["log10_p"], "0.00")
    put(ws, row, 11, r["p"], "0.00E+00")
    put(ws, row, 12, r["q"], "0.00E+00")
    put(ws, row, 13, '=IF(L{0}<0.05,"Yes","No")'.format(row))
last = r0 + len(fig2_rows)
srow = last + 2
put(ws, srow, 1, "Significant, k = 3", font=F_LABEL)
put(ws, srow, 4, '=COUNTIFS(C{0}:C{1},3,M{0}:M{1},"Yes")'.format(r0 + 1, last))
put(ws, srow + 1, 1, "Significant, k = 4", font=F_LABEL)
put(ws, srow + 1, 4, '=COUNTIFS(C{0}:C{1},4,M{0}:M{1},"Yes")'.format(r0 + 1, last))
put(ws, srow + 2, 1, "Under-represented (Z < 0)", font=F_LABEL)
put(ws, srow + 2, 4, "=COUNTIF(I{0}:I{1},\"<0\")".format(r0 + 1, last))
put(ws, srow + 3, 1, "Smallest Z", font=F_LABEL)
put(ws, srow + 3, 4, "=MIN(I{0}:I{1})".format(r0 + 1, last), "0.000")

# ── Fig 3 ─────────────────────────────────────────────────────────────────
ws = wb.create_sheet("Fig 3")
r0 = title_block(ws, "Fig. 3. Five-family TF-regulates-TF cascades (all 21 combinations)", [
    "Edge A → B: at least one TF in family A has a π4 binding site in the promoter of a TF that is a "
    "member of family B. Weight = number of such (TF in A, TF in B) pairs.",
    "Mutual: both A → B and B → A exist (vermillion in Fig. 3); one-way: only A → B (blue). "
    "Combination # is the itertools.combinations order of F1-F7, as printed on the panels.",
    "Topology group: combinations with the same (isolated nodes, directed edges, mutual pairs) "
    "share a group (labeled under each panel of Fig. 3).",
])
put(ws, r0, 1, "Summary", font=F_LABEL)
sum_hdr = r0 + 1
header(ws, sum_hdr, ["Combination #", "Families", "Directed edges", "Mutual pairs",
                     "Isolated nodes", "Topology group"], freeze=False)
edge_start = sum_hdr + len(fig3_combos) + 4
edge_hdr = edge_start
n_edges_total = sum(c["G"].number_of_edges() for c in fig3_combos)
e_first, e_last = edge_hdr + 1, edge_hdr + n_edges_total
for i, c in enumerate(fig3_combos):
    row = sum_hdr + 1 + i
    put(ws, row, 1, c["num"])
    put(ws, row, 2, c["label"])
    put(ws, row, 3, "=COUNTIF($A${0}:$A${1},A{2})".format(e_first, e_last, row))
    put(ws, row, 4, '=COUNTIFS($A${0}:$A${1},A{2},$E${0}:$E${1},"mutual")/2'.format(
        e_first, e_last, row))
    put(ws, row, 5, c["n_isolated"])
    put(ws, row, 6, c["group"])
put(ws, edge_start - 1, 1, "Edge list", font=F_LABEL)
header(ws, edge_hdr, ["Combination #", "Families", "Source family", "Target family",
                      "Edge type", "Weight (TF pairs)"], [14, 22, 13, 13, 11, 16], freeze=False)
row = edge_hdr + 1
for c in fig3_combos:
    G = c["G"]
    for u, v, d in sorted(G.edges(data=True)):
        put(ws, row, 1, c["num"])
        put(ws, row, 2, c["label"])
        put(ws, row, 3, "F%d" % u)
        put(ws, row, 4, "F%d" % v)
        put(ws, row, 5, "mutual" if G.has_edge(v, u) else "one-way")
        put(ws, row, 6, d["weight"])
        row += 1
assert row - 1 == e_last

# ── Fig 4 ─────────────────────────────────────────────────────────────────
ws = wb.create_sheet("Fig 4")
r0 = title_block(ws, "Fig. 4. Regulatory-link maximum-parsimony tree", [
    "Characters: one per π4-covered target gene; a family scores 1 if any member TF binds that target. "
    "Exhaustive search over all 945 unrooted binary trees, Fitch parsimony, characters weighted by "
    "identical-column count.",
    "Rooted at the midpoint of F5's branch (RNA-World outgroup proxy). Branch lengths = weighted "
    "character changes (ACCTRAN-style Fitch down-pass).",
])
info = [
    ("Parsimony score (changes)", fig4_score, "#,##0"),
    ("Equally parsimonious trees", n_tied, "0"),
    ("Newick (rooted, branch lengths)", newick, None),
    ("Characters (target genes)", int(raw.shape[1]), "#,##0"),
    ("Constant characters (dropped from search)", int(g4["n_constant"]), "#,##0"),
    ("Parsimony-informative characters", int(g4["n_informative"]), "#,##0"),
    ("Single-family (autapomorphic) characters", int(g4["n_autapomorphy"]), "#,##0"),
    ("Note", "F7 (HMRA1) has no π4-covered member TF, so its row is all 0 and its placement is "
             "not supported by F7-specific evidence.", None),
]
for i, (a, b, fmt) in enumerate(info):
    put(ws, r0 + i, 1, a, font=F_LABEL)
    c = put(ws, r0 + i, 3, b, fmt)
    if isinstance(b, str):
        c.alignment = Alignment(wrap_text=False)
mat_hdr = r0 + len(info) + 2
put(ws, mat_hdr - 1, 1, "Binary character matrix (1 = family regulates target)", font=F_LABEL)
header(ws, mat_hdr, ["Target gene ID", "Target gene name"] + ["F%d" % t for t in TAXA]
       + ["Character class"], [16, 16, 6, 6, 6, 6, 6, 6, 6, 18])
for j, tgt in enumerate(all_targets):
    row = mat_hdr + 1 + j
    col = raw[:, j]
    ones = int(col.sum())
    cls = ("constant" if ones in (0, len(TAXA)) else
           "autapomorphic" if ones in (1, len(TAXA) - 1) else "informative")
    put(ws, row, 1, tgt)
    name = _pi4_names.get(tgt)
    put(ws, row, 2, name if isinstance(name, str) else None)
    for i_t in range(len(TAXA)):
        put(ws, row, 3 + i_t, int(col[i_t]))
    put(ws, row, 3 + len(TAXA), cls)

# ── Pi estimates ──────────────────────────────────────────────────────────
ws = wb.create_sheet("Pi estimates")
r0 = title_block(ws, "Per-family inheritance probability π from the seven estimators", [
    "Per-TF estimates are averaged over each family's member TFs that have a value "
    "(the two-level mean of Method 1, Eq. 1). n = member TFs contributing.",
    "M1 used in the Fig. 2 tests = family evidence score eᵢ normalized by Σ eⱼ over all 7 families. "
    "M1 π(1) = eᵢ plus DNA-binding (+0.05) and activator (+0.03) corrections, capped at 1.",
    "M2 (Theorem 4 inversion) identifies only π̂ = Σπᵢ for a given motif, not per-family πᵢ; it is "
    "reported per motif in the second table, allocated across families by eᵢ.",
    "M3 is the YFL039C cross-strain mean, applied to every family. M5 = π3 TFBS retention, "
    "M6 = π2 mean pairwise protein identity, M7 = π4 = 1 − IC-weighted binding-site polymorphism. "
    "Ensemble = mean of M1 π(1), M3-M7.",
])
pcols = ["Family", "GO ID", "Size (sᵢ)", "Evidence score eᵢ", "M1 (normalized, used in Fig. 2)",
         "M1 π(1)", "M2", "M3 SNP divergence", "M4 Consensus-adjusted", "M5 π3 TFBS conservation",
         "M6 π2 Sequence homology", "M7 π4 Binding-site SNPs", "Ensemble (M1 π(1), M3-M7)",
         "n TFs: M1/M4", "n TFs: M5", "n TFs: M6", "n TFs: M7"]
header(ws, r0, pcols, [8, 13, 9, 11, 14, 10, 18, 12, 13, 14, 14, 14, 15, 9, 9, 9, 9], freeze=False)
for i, pr in enumerate(pi_rows):
    row = r0 + 1 + i
    put(ws, row, 1, pr["fam"])
    put(ws, row, 2, pr["go"])
    put(ws, row, 3, pr["size"])
    put(ws, row, 4, pr["ev"], "0.0000")
    put(ws, row, 5, "=D{0}/SUM($D${1}:$D${2})".format(row, r0 + 1, r0 + len(pi_rows)), "0.0000")
    put(ws, row, 6, pr["m1"], "0.0000")
    put(ws, row, 7, "per motif only (see below)")
    for col, key in ((8, "m3"), (9, "m4"), (10, "m5"), (11, "m6"), (12, "m7")):
        put(ws, row, col, pr[key] if pr[key] is not None else "no data", "0.0000")
    put(ws, row, 13, "=AVERAGE(F{0},H{0}:L{0})".format(row), "0.0000")
    n1, n4, n5, n6, n7 = pr["n"]
    put(ws, row, 14, n1)
    put(ws, row, 15, n5)
    put(ws, row, 16, n6)
    put(ws, row, 17, n7)

m2_hdr = r0 + len(pi_rows) + 4
put(ws, m2_hdr - 1, 1, "Method 2: π̂ recovered per motif by inverting Theorem 4 at the Fig. 2 "
                       "observed count, allocated across families ∝ eᵢ", font=F_LABEL)
header(ws, m2_hdr, ["Combination", "Observed count", "π̂ (M2)"] + ["F%d" % t for t in TAXA], freeze=False)
for i, (label, obs, ph, alloc) in enumerate(m2_rows):
    row = m2_hdr + 1 + i
    put(ws, row, 1, label)
    put(ws, row, 2, obs, "#,##0")
    put(ws, row, 3, ph if ph is not None else "no solution", "0.0000")
    for j, t in enumerate(TAXA):
        v = alloc.get("F%d" % t)
        put(ws, row, 4 + j, v, "0.0000")

OUT_PATH.parent.mkdir(exist_ok=True)
wb.save(OUT_PATH)

print("\nSaved %s" % OUT_PATH)
print("Fig 2: %d tests, %d significant, min Z = %.4f (%s)" % (
    n_tests, int((q < FDR_ALPHA).sum()), zs.min(), fig2_rows[int(zs.argmin())]["label"]))
print("Fig 3: %d combos, %d edges, %d topology groups" % (
    len(fig3_combos), n_edges_total, len({c["group"] for c in fig3_combos})))
for c in (fig3_combos[0], fig3_combos[-1]):
    print("  #%d %s: %d edges, %d mutual" % (c["num"], c["label"], c["n_edges"], c["n_mutual"]))
print("Fig 4: score %.0f, %d tied tree(s), %d characters" % (fig4_score, n_tied, raw.shape[1]))
print("  Newick: %s" % newick)
print("Pi:")
for pr in pi_rows:
    print("  ", {k: (round(v, 4) if isinstance(v, float) else v) for k, v in pr.items()})
